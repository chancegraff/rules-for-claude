#!/usr/bin/env python3
"""fit_titles.py — the sanctioned ONE-LINE TITLE auto-fit.

Attentive content-header titles read best "short and sweet" on a SINGLE line.
The canonical header box is narrow (~2.55in), so a header that is even slightly
long wraps to 2-4 lines and — on the icon / stat / card layouts — the wrapped
lines spill down into the row beneath the title. This module performs the ONE
geometry change we permit on a duplicated canonical slide: it WIDENS a
content-header box just enough to hold its longest line on one line.

It is deliberately the narrowest possible exception to the "never resize a
duplicated slide" rule:

  * only CONTENT HEADERS — the mid-size (~16-28pt) title/subtitle block pinned to
    the slide's top-left corner. The big display titles on covers, section
    breakers and statements (>= 28pt), the tiny eyebrow (< 16pt), the stat
    numbers lower down the slide, and multi-line agenda lists are all left
    EXACTLY as designed;
  * only ever GROWS the box, never shrinks it;
  * only when the result is a genuine ONE-LINE header — if the copy is so long it
    could not fit one line even using all the free width, the box is left alone
    (and the linter tells the author to shorten it);
  * never past a neighbouring shape (leaves a gutter) or the slide's right edge;
  * never touches x / y / height / font / colour / z-order or any other element.

Because it only grows a box into empty space to fit text that is already there,
it can never make a slide worse. Widths are computed with the SemiBold
char-width model (Attentive headers render SemiBold), so a header that fits here
also fits in Google Slides.

Public API
----------
    import fit_titles
    changed = fit_titles.widen_titles(slide, slide_w_in)   # python-pptx slide
    # returns [(shape_name, old_width_in, new_width_in), ...]  (widened boxes)

    # geometry helper reused by the linter so its fit check matches the build:
    fit_titles.available_width(slide, shp, slide_w_in)
    fit_titles.is_content_header(shp)
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx.oxml.ns import qn
from pptx.util import Emu

# fit_text lives in the sibling formatting skill; import it the way the other
# builder modules do. Degrade to a no-op if unavailable — never crash a build
# over an optional cosmetic pass.
_FIT_DIR = Path(__file__).parent.parent / "attentive-slide-formatting" / "scripts"
if str(_FIT_DIR) not in sys.path:
    sys.path.insert(0, str(_FIT_DIR))
try:
    import fit_text
except Exception:  # pragma: no cover - defensive
    fit_text = None

EMU_PER_IN = 914400

# A content header is the mid-size title/subtitle block in the top-left corner.
CONTENT_MIN_PT = 16.0        # below this = eyebrow / label / cover subtitle
DISPLAY_TITLE_MIN_PT = 28.0  # at/above this = big display title (cover/section)
HEADER_TOP_MAX_IN = 1.30     # header sits near the top of the slide
HEADER_LEFT_MAX_IN = 2.00    # header is left-aligned
MAX_SEGMENTS = 2             # title + optional subtitle (excludes agenda lists)

# Attentive content headers render SemiBold — size the box against that weight so
# a one-line fit here holds in Google Slides too.
HEADER_WEIGHT_FACE = "Inter SemiBold"

GUTTER_IN = 0.15             # clear space kept before a right-neighbour
RIGHT_MARGIN_IN = 0.20       # clear space kept before the slide's right edge
PAD_IN = 0.0375             # per-side text inset (matches fit_text's default)


def _emu_in(v):
    return Emu(v).inches if v is not None else None


def ph_type(shp):
    el = getattr(shp, "_element", None)
    if el is None:
        return None
    ph = el.find(".//" + qn("p:ph"))
    return ph.get("type") if ph is not None else None


def _size_pt(shp):
    """First explicit run/paragraph font size (pt) in the shape, or None."""
    body = shp._element.find(".//" + qn("p:txBody"))
    if body is None:
        return None
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for el in body.findall(".//" + qn(tag)):
            sz = el.get("sz")
            if sz:
                try:
                    return int(sz) / 100.0
                except ValueError:
                    return None
    return None


def title_segments(shp) -> list:
    """The hard lines a header must each fit on: every paragraph, further split on
    explicit ``<a:br/>`` breaks (the header stacks title + subtitle as two runs
    separated by a break, so each must fit on its own line)."""
    if not shp.has_text_frame:
        return []
    body = shp._element.find(".//" + qn("p:txBody"))
    if body is None:
        return []
    segs = []
    for p in body.findall(qn("a:p")):
        cur = ""
        for child in p:
            if child.tag == qn("a:r"):
                t = child.find(qn("a:t"))
                cur += (t.text or "") if t is not None else ""
            elif child.tag == qn("a:br"):
                segs.append(cur)
                cur = ""
        segs.append(cur)
    return [s for s in segs if s.strip()]


def is_content_header(shp) -> bool:
    """True for the mid-size title/subtitle block pinned to the top-left — the
    only thing this module ever widens. Identified by geometry + size + line
    count so it catches both the title-placeholder headers and the plain-textbox
    headers the stat layouts use, while excluding display titles, eyebrows, stat
    numbers, and agenda lists."""
    try:
        if not shp.has_text_frame:
            return False
        size = _size_pt(shp)
        if size is None or not (CONTENT_MIN_PT <= size < DISPLAY_TITLE_MIN_PT):
            return False
        top = _emu_in(shp.top)
        left = _emu_in(shp.left)
        if top is None or left is None:
            return False
        if top > HEADER_TOP_MAX_IN or left > HEADER_LEFT_MAX_IN:
            return False
        segs = title_segments(shp)
        return 1 <= len(segs) <= MAX_SEGMENTS
    except Exception:
        return False


def available_width(slide, shp, slide_w_in: float) -> float:
    """The widest the header box may grow (inches) before it would cross a
    neighbouring shape or the slide's right margin. Only shapes to the RIGHT of
    the header that vertically overlap its top strip constrain it; content lower
    down the slide (columns, stat cards, icon rows) does not, because a one-line
    header never reaches it."""
    left = _emu_in(shp.left) or 0.0
    top = _emu_in(shp.top) or 0.0
    h = _emu_in(shp.height) or 0.9
    band_lo, band_hi = top, top + max(h, 0.9)
    cap_right = slide_w_in - RIGHT_MARGIN_IN
    for other in slide.shapes:
        if other._element is shp._element:
            continue
        try:
            ol = _emu_in(other.left)
            ot = _emu_in(other.top)
            ow = _emu_in(other.width)
            oh = _emu_in(other.height)
        except Exception:
            continue
        if None in (ol, ot, ow, oh):
            continue
        if ol <= left + 0.01:                        # not to the right
            continue
        if ot < band_hi and (ot + oh) > band_lo:     # overlaps the top strip
            cap_right = min(cap_right, ol - GUTTER_IN)
    return max(_emu_in(shp.width) or 0.0, cap_right - left)


def required_width(segments, size_pt: float):
    """Box width (inches) for the longest segment to sit on ONE line, using the
    SemiBold char-width model. None if unknown."""
    if fit_text is None or not size_pt:
        return None
    cpi, _ = fit_text.chars_per_inch(HEADER_WEIGHT_FACE, size_pt)
    if not cpi:
        return None
    longest = max((len(s) for s in segments), default=0)
    if longest == 0:
        return None
    return longest / cpi + 2 * PAD_IN


def widen_titles(slide, slide_w_in: float) -> list:
    """Widen every content header on ``slide`` toward one line, within free space.
    Returns ``[(name, old_in, new_in), ...]`` for boxes that changed."""
    if fit_text is None:
        return []
    changed = []
    for shp in list(slide.shapes):
        if not is_content_header(shp):
            continue
        req = required_width(title_segments(shp), _size_pt(shp))
        if not req:
            continue
        cur = _emu_in(shp.width) or 0.0
        if req <= cur + 0.02:
            continue                                 # already fits one line
        avail = available_width(slide, shp, slide_w_in)
        if req > avail:
            continue                                 # can't reach one line here
        shp.width = Emu(int(round(req * EMU_PER_IN)))
        changed.append((shp.name, round(cur, 2), round(req, 2)))
    return changed


__all__ = [
    "widen_titles", "available_width", "required_width", "title_segments",
    "is_content_header", "ph_type", "DISPLAY_TITLE_MIN_PT", "CONTENT_MIN_PT",
]
