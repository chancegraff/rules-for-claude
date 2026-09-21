#!/usr/bin/env python3
"""slide_facts.py — derive per-slide facts for the 164 canonical template slides.

WHY THIS EXISTS
---------------
Rebranding the AI Guild deck turned up a set of traps in the canonical template
that no existing data file records. Every one was found by hand, mid-build:

  * canonical 149 is the RED-X (negative) twin of 150's green-check layout, so
    filling it with three positive pillars silently inverted their meaning
  * canonical 142's last two bullet rows carry a roadmap "QX" badge, which
    branded two personality traits as upcoming features
  * canonical 152's image slots ship WITH stock photos (a bedroom, sunscreen, a
    man in sunglasses) — leave them and they ship; replace them and you inherit
    their crop
  * those pre-set crops (152's third slot cuts 68% off the bottom) survive an
    image swap and mangle any replacement
  * multi-slot slides are addressed in SHAPE-ID order, which is NOT left-to-right
    order — 150's columns run 3303 / 3311 / 3307 and 152's image slots run
    3338 / 3337 / 3341, so writing content in id order shuffles it on screen
  * per-slot capacity varies wildly and is only hinted at inside the dummy copy
    ("~20 character point 1", "<40 char sub-point"); 92's stat captions cap
    around 10 characters

`inventory.json` records placeholder TEXT and `slide_scores.json` records
density/impact SCORES. Neither records any of the above, so each rebrand
rediscovers them by rendering and squinting.

Everything here is DERIVED from the .pptx, not hand-typed, so it stays correct
when the template changes — re-run it. Nothing is written back to the template;
this only reads.

    python slide_facts.py                 # print summary; do not modify the skill
    python slide_facts.py --slide 150     # explain one slide
    python slide_facts.py --traps         # only slides carrying a trap
    python slide_facts.py --write /path/to/slide_facts.json  # explicit refresh
"""
from __future__ import annotations

import argparse
import collections
import json
import os
import re
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

try:
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:  # pragma: no cover
    MSO_SHAPE_TYPE = None

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_NAME = "Template—2026 NEW Attentive Company Deck Template.pptx"


def _resolve_template(name=TEMPLATE_NAME):
    """Find the packaged canonical deck next to this script or one level up."""
    requested = os.path.expanduser(name)
    if os.path.isabs(requested):
        candidates = [requested]
    elif name == TEMPLATE_NAME:
        candidates = [os.path.join(HERE, name), os.path.join(HERE, "..", name), name]
    else:
        candidates = [requested, os.path.join(HERE, name), os.path.join(HERE, "..", name)]
    for cand in candidates:
        if os.path.isfile(cand):
            return cand
    return os.path.join(HERE, name)  # nonexistent: let the caller raise clearly


DEFAULT_TEMPLATE = _resolve_template()
OUT_PATH = os.path.join(HERE, "slide_facts.json")
EMU_PER_IN = 914400.0

# Capacity hints the template authors embedded in their own dummy copy.
_CAP_PATTERNS = [
    (re.compile(r"(\d+)\s*-\s*(\d+)\s*character", re.I), "range"),
    (re.compile(r"[<~]\s*(\d+)\s*char", re.I), "single"),
    (re.compile(r"roughly\s*(\d+)\s*characters", re.I), "single"),
    (re.compile(r"less than\s*(\d+)\s*characters", re.I), "single"),
    (re.compile(r"contains\s*~?(\d+)\s*characters", re.I), "single"),
]

NEGATIVE_GLYPHS = {"✕", "✗", "✘", "×", "x"}
POSITIVE_GLYPHS = {"✓", "✔"}


def _in(v):
    return round((v or 0) / EMU_PER_IN, 3)


def _char_cap(text: str):
    """Pull the author's own capacity hint out of a dummy string."""
    for pat, kind in _CAP_PATTERNS:
        m = pat.search(text)
        if m:
            if kind == "range":
                return int(m.group(2))
            return int(m.group(1))
    return None


def _font_pt(shape, slide):
    """Explicit run/paragraph size, else the layout placeholder's, else None."""
    if not getattr(shape, "has_text_frame", False):
        return None
    sizes = []
    for para in shape.text_frame.paragraphs:
        if para.font.size is not None:
            sizes.append(para.font.size.pt)
        for run in para.runs:
            if run.font.size is not None:
                sizes.append(run.font.size.pt)
    if sizes:
        return max(sizes)
    try:
        if shape.is_placeholder:
            idx = shape.placeholder_format.idx
            for lph in slide.slide_layout.placeholders:
                if lph.placeholder_format.idx == idx:
                    for para in lph.text_frame.paragraphs:
                        if para.font.size is not None:
                            return para.font.size.pt
                        for run in para.runs:
                            if run.font.size is not None:
                                return run.font.size.pt
    except (AttributeError, KeyError, ValueError):
        pass
    return None


def _est_char_cap(width_in, font_pt):
    """Rough single-line capacity when the author left no hint in the dummy copy.

    Slide 92's stat captions are the case that motivated this: the dummy text is
    just "Lorem ipsum " with no character hint, yet the box only holds about ten
    characters, so "weekly Guild sync" overflowed with nothing to warn me.
    Inter averages ~0.5em per character, so chars ≈ width / (pt * 0.5 / 72).
    """
    if not width_in or not font_pt:
        return None
    return max(1, int(width_in / (font_pt * 0.5 / 72.0)))


def _bullet_glyphs(shape):
    out = set()
    if not getattr(shape, "has_text_frame", False):
        return out
    for el in shape.text_frame._txBody.iter():
        if el.tag == qn("a:buChar"):
            ch = el.get("char")
            if ch:
                out.add(ch)
    return out


def _fill_hex(shape):
    try:
        if shape.fill.type is not None:
            return str(shape.fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError):
        pass
    return None


def _kind(shape):
    if getattr(shape, "has_chart", False):
        return "chart"
    if getattr(shape, "has_table", False):
        return "table"
    if MSO_SHAPE_TYPE is not None and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "picture_filled"
    try:
        if shape.is_placeholder and "PICTURE" in str(shape.placeholder_format.type):
            return "picture_empty"
    except (AttributeError, ValueError):
        pass
    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        return "text"
    return "deco"


def analyse_slide(slide, number):
    text_slots, image_slots, accents = [], [], []
    glyphs = set()
    tables = charts = 0

    for shp in slide.shapes:
        kind = _kind(shp)
        x, y = _in(shp.left), _in(shp.top)
        w, h = _in(shp.width), _in(shp.height)
        if kind == "table":
            tables += 1
        elif kind == "chart":
            charts += 1
        elif kind in ("picture_filled", "picture_empty"):
            crop = {}
            for side in ("left", "right", "top", "bottom"):
                try:
                    crop[side] = round(getattr(shp, f"crop_{side}"), 4)
                except (AttributeError, TypeError):
                    crop[side] = 0.0
            if w * h < 0.6:          # tiny marks (logos, dots) aren't content slots
                continue
            image_slots.append({
                "shape": shp.name, "x": x, "y": y, "w": w, "h": h,
                "aspect": round(w / h, 3) if h else None,
                "ships_filled": kind == "picture_filled",
                "preset_crop": crop,
                "has_preset_crop": any(v > 0.001 for v in crop.values()),
            })
        elif kind == "text":
            txt = shp.text_frame.text
            runs = [r for p in shp.text_frame.paragraphs for r in p.runs]
            paras = [p.text for p in shp.text_frame.paragraphs]
            glyphs |= _bullet_glyphs(shp)
            hinted = _char_cap(txt)
            font_pt = _font_pt(shp, slide)
            text_slots.append({
                "shape": shp.name, "x": x, "y": y, "w": w, "h": h,
                "text": txt,
                "n_runs": len(runs),
                "n_paras": len([p for p in paras if p.strip()]),
                "font_pt": font_pt,
                "char_cap": hinted,
                # geometric fallback for slots the authors left un-hinted
                "est_char_cap": None if hinted else _est_char_cap(w, font_pt),
            })
        else:
            if w < 0.7 and h < 0.7:
                hexf = _fill_hex(shp)
                if hexf:
                    accents.append(hexf)

    # left-to-right, then top-to-bottom — the order a reader consumes them, and
    # the order that shape-id sorting does NOT give you
    reading = lambda s: (round(s["y"], 1), round(s["x"], 1))
    text_slots.sort(key=lambda s: (s["x"], s["y"]))
    image_slots.sort(key=lambda s: (s["x"], s["y"]))

    # repeated slots = identical dummy copy appearing more than once. That count
    # is the slide's true arity (how many cards / columns / stats it holds).
    # The length floor is deliberately low (>4, not >20) so short repeated labels
    # count too — slide 92's three stat captions are all just "Lorem ipsum " and
    # were invisible at the higher threshold. Guard against grouping unrelated
    # short labels by also requiring a shared width.
    counts = collections.Counter(s["text"] for s in text_slots if len(s["text"]) > 4)
    groups = []
    for txt, n in counts.items():
        if n < 2:
            continue
        members = [s for s in text_slots if s["text"] == txt]
        widths = [m["w"] for m in members]
        if max(widths) - min(widths) > 0.35:      # not a real repeated group
            continue
        first = members[0]
        groups.append({
            "arity": n,
            "reading_order": [m["shape"] for m in members],   # already x-sorted
            "n_runs": first["n_runs"],
            "char_cap": _char_cap(txt) or first["est_char_cap"],
            "cap_is_estimated": _char_cap(txt) is None,
            "sample": txt[:70],
        })
    groups.sort(key=lambda g: -g["arity"])

    all_text = " ".join(s["text"] for s in text_slots)
    polarity = "neutral"
    if glyphs & NEGATIVE_GLYPHS:
        polarity = "negative"
    elif glyphs & POSITIVE_GLYPHS:
        polarity = "positive"
    if re.search(r"negative points|problem \d|things you said", all_text, re.I):
        polarity = "negative"
    elif re.search(r"positive points", all_text, re.I):
        polarity = "positive"

    return {
        "number": number,
        "text_slots": text_slots,
        "image_slots": image_slots,
        "repeated_groups": groups,
        "max_arity": groups[0]["arity"] if groups else 0,
        "bullet_glyphs": sorted(glyphs),
        "polarity": polarity,
        "accent_fills": sorted(set(accents)),
        "n_tables": tables,
        "n_charts": charts,
        "has_roadmap_badge": bool(re.search(r"upcoming feature", all_text, re.I)),
    }


def traps_for(f):
    """Human-readable warnings — the things that cost time this session."""
    out = []
    if f["polarity"] == "negative":
        out.append(f"NEGATIVE variant (glyph {'/'.join(f['bullet_glyphs']) or '—'}) "
                   f"— do not use for positive content")
    if f["has_roadmap_badge"]:
        out.append("carries a roadmap 'QX' badge on its last rows — not a plain list")
    filled = [s for s in f["image_slots"] if s["ships_filled"]]
    if filled:
        out.append(f"{len(filled)} image slot(s) SHIP WITH stock photos — they "
                   f"render unless replaced")
    cropped = [s for s in f["image_slots"] if s["has_preset_crop"]]
    if cropped:
        worst = max(cropped, key=lambda s: max(s["preset_crop"].values()))
        out.append(f"{len(cropped)} image slot(s) carry a pre-set crop "
                   f"(up to {int(max(worst['preset_crop'].values())*100)}%) — "
                   f"reset it when swapping images")
    empty = [s for s in f["image_slots"] if not s["ships_filled"]]
    if empty:
        out.append(f"{len(empty)} image slot(s) ship EMPTY — they render as grey "
                   f"boxes unless filled")
    for g in f["repeated_groups"]:
        ids = [int(m.split(";")[1]) for m in g["reading_order"] if ";" in m]
        if ids and ids != sorted(ids):
            out.append(f"{g['arity']}-slot group is NOT in shape-id order — "
                       f"reading order is {', '.join(g['reading_order'])}")
    tight = [s for s in f["text_slots"]
             if (s["char_cap"] or s.get("est_char_cap"))
             and (s["char_cap"] or s["est_char_cap"]) <= 12]
    if tight:
        names = ", ".join(s["shape"].split(";")[1] for s in tight[:4])
        out.append(f"{len(tight)} very tight slot(s), cap <= 12 chars ({names})")
    return out


def build(template=DEFAULT_TEMPLATE):
    prs = Presentation(template)
    facts = {}
    for i, slide in enumerate(prs.slides, start=1):
        f = analyse_slide(slide, i)
        f["traps"] = traps_for(f)
        facts[str(i)] = f
    return {
        "template": os.path.basename(template),
        "slide_count": len(facts),
        "generated_by": "slide_facts.py — derived, not hand-written; re-run after "
                        "any template change",
        "slides": facts,
    }


def print_slide(f):
    print(f"\ncanonical slide {f['number']}")
    print(f"  polarity {f['polarity']}  glyphs {f['bullet_glyphs']}  "
          f"max arity {f['max_arity']}")
    for g in f["repeated_groups"]:
        print(f"  {g['arity']}-slot group · {g['n_runs']} runs · cap {g['char_cap']}")
        print(f"      reading order: {', '.join(g['reading_order'])}")
    for s in f["image_slots"]:
        print(f"  image {s['shape']} {s['w']}x{s['h']} ar{s['aspect']} "
              f"{'FILLED' if s['ships_filled'] else 'empty'}"
              f"{' preset-crop' if s['has_preset_crop'] else ''}")
    tight = [(s["shape"], s["char_cap"]) for s in f["text_slots"] if s["char_cap"]]
    if tight:
        print("  char caps: " + ", ".join(f"{n.split(';')[1]}={c}" for n, c in tight))
    for t in f["traps"]:
        print(f"  ⚠ {t}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--slide", type=int)
    ap.add_argument("--traps", action="store_true")
    ap.add_argument("--template", default=DEFAULT_TEMPLATE)
    ap.add_argument("--write", metavar="PATH",
                    help="write regenerated facts to this explicit output path")
    args = ap.parse_args()

    data = build(args.template)
    if args.slide:
        print_slide(data["slides"][str(args.slide)])
        return
    n_traps = sum(1 for f in data["slides"].values() if f["traps"])
    if args.write:
        out_path = os.path.abspath(os.path.expanduser(args.write))
        with open(out_path, "w") as fh:
            json.dump(data, fh, indent=1)
        print(f"wrote {out_path} — {data['slide_count']} slides, "
              f"{n_traps} carry at least one trap")
    else:
        print(f"inspected {data['slide_count']} slides; "
              f"{n_traps} carry at least one trap")
    if args.traps:
        for num, f in sorted(data["slides"].items(), key=lambda x: int(x[0])):
            if f["traps"]:
                print(f"\n {num}:")
                for t in f["traps"]:
                    print(f"    ⚠ {t}")
    else:
        counts = collections.Counter(
            t.split("—")[0].split("(")[0].strip()
            for f in data["slides"].values() for t in f["traps"])
        print("\ntrap frequency across the template:")
        for t, c in counts.most_common():
            print(f"  {c:>3}×  {t}")


if __name__ == "__main__":
    main()
