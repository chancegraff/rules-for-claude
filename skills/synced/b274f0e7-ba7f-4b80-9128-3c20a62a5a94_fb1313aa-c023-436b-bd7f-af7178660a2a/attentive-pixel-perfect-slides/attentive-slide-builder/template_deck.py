#!/usr/bin/env python3
"""template_deck.py — the canonical-template DUPLICATION engine.

This is the sanctioned way to produce Attentive slides. It does the opposite of
a generator: it never builds geometry from coordinates. It opens the canonical
Attentive template deck, DUPLICATES the exact slides you ask for (preserving all
geometry, spacing, fonts, colors, dividers, backgrounds, image frames, tables,
icons, shapes, and z-order byte-for-byte), and replaces ONLY the placeholder
text and images.

Core idea
---------
We work *inside a copy of the canonical deck*, so every slide master, layout,
theme, font reference, and media file is already present and correctly linked.
For each slide you want, we clone the matching canonical slide within that copy
(rels and media resolve because we stay in the same package), fill its
placeholders, and at the end drop the original canonical slides so only your
filled clones remain, in the order you specified.

The geometry is therefore identical to the source slide — there is no rebuild,
no eyeballing, no from-scratch rendering. If a clone looks wrong, the source
slide looked that way too.

Public API
----------
    deck = TemplateDeck("Template—2026 NEW Attentive Company Deck Template.pptx")
    deck.list_slide_text(24)                 # inspect a template slide's text
    s = deck.use(24)                         # duplicate canonical slide 24
    s.replace_text({"Old header": "New header", "Old bullet": "New bullet"})
    s.replace_image("Picture 3", "assets/hero.png")
    deck.save("out.pptx")                    # writes only the duplicated slides

Or drive it from a plan JSON via build_from_plan.py.
"""
from __future__ import annotations

import copy
import os
import re
import shutil
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn


def _iter_shapes(shapes):
    """Yield every shape INCLUDING those nested inside groups. Text lives inside
    grouped shapes on several template slides (e.g. timeline 162); without
    recursing, those placeholders are never found or filled."""
    for shp in shapes:
        yield shp
        if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shp.shapes)

_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_SHAPE_TAGS = ("p:sp", "p:pic", "p:graphicFrame", "p:grpSp", "p:cxnSp")
_SHAPE_TAG_SET = {qn(t) for t in _SHAPE_TAGS}


def _remap_rels(src_part, new_part, root_el):
    """Re-create, in new_part, every relationship referenced inside root_el and
    rewrite the r:id / r:embed / r:link attributes to the new rIds. This is what
    lets a cloned slide keep pointing at the right images/media without
    colliding with relationships the new slide part already owns."""
    old_to_new: dict[str, str] = {}
    for el in root_el.iter():
        for name, val in list(el.attrib.items()):
            if not name.startswith("{%s}" % _REL_NS):
                continue
            old_rid = val
            if old_rid in old_to_new:
                el.set(name, old_to_new[old_rid])
                continue
            if old_rid not in src_part.rels:
                continue
            rel = src_part.rels[old_rid]
            if rel.is_external:
                new_rid = new_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                new_rid = new_part.relate_to(rel.target_part, rel.reltype)
            old_to_new[old_rid] = new_rid
            el.set(name, new_rid)


# Default emphasis colors for the progress-divider active/muted sections.
_DIVIDER_ACTIVE = RGBColor(0x1E, 0x1C, 0x1C)   # ink
_DIVIDER_MUTED = RGBColor(0xBC, 0xB8, 0xAC)    # warm gray
_SECTION_RE = re.compile(r"^\s*Section\s*0?\d+\s*$", re.I)


def _replace_in_shape(shp, replacements: dict[str, str]) -> set:
    """Run-level substring replacement scoped to ONE shape (preserves each
    run's formatting); falls back to paragraph-level for keys that span runs.
    Because it is scoped to a single shape, identical placeholder strings that
    repeat across sibling shapes (columns, stat cards, table cells) can each be
    filled with a DIFFERENT value — which whole-slide find/replace cannot do.
    Returns the set of keys applied within this shape."""
    applied: set = set()
    if not shp.has_text_frame:
        return applied
    for para in shp.text_frame.paragraphs:
        for run in para.runs:
            for find, repl in replacements.items():
                if find in run.text:
                    run.text = run.text.replace(find, repl)
                    applied.add(find)
        para_text = "".join(r.text for r in para.runs)
        for find, repl in replacements.items():
            if find in applied:
                continue
            if find and find in para_text and para.runs:
                para.runs[0].text = para_text.replace(find, repl)
                for extra in para.runs[1:]:
                    extra.text = ""
                applied.add(find)
    return applied


def _set_cell_text(cell, text: str):
    """Set a table cell's text while preserving its first run's formatting
    (fill/borders/paragraph props are the cell's own and are untouched)."""
    tf = cell.text_frame
    para = tf.paragraphs[0]
    if para.runs:
        para.runs[0].text = text
        for extra in para.runs[1:]:
            extra.text = ""
    else:
        cell.text = text


class Slide:
    """A duplicated canonical slide. Only text/image replacement is exposed —
    there is deliberately no API to move, resize, or restyle elements."""

    def __init__(self, slide, source_number: int):
        self._slide = slide
        self.source_number = source_number

    # --- inspection -------------------------------------------------------
    def text_shapes(self):
        out = []
        for shp in _iter_shapes(self._slide.shapes):
            if shp.has_text_frame:
                out.append((shp.name, shp.text_frame.text))
        return out

    def find_shape(self, name: str):
        for shp in _iter_shapes(self._slide.shapes):
            if shp.name == name:
                return shp
        return None

    # --- text replacement (formatting preserved) --------------------------
    def replace_text(self, replacements: dict[str, str]) -> list[str]:
        """Replace placeholder copy across the WHOLE slide. `replacements` maps
        existing text -> new text. Matching is tried run-by-run first (preserves
        that run's exact formatting); if a key spans multiple runs it is matched
        against the whole paragraph and rewritten into the first run. Returns the
        list of keys applied.

        Note: a key that repeats across sibling shapes (columns, stat cards) is
        replaced EVERYWHERE with the same value. To fill such repeated
        placeholders with different values, use `fill_shapes` / `fill_table`."""
        applied: set[str] = set()
        for shp in _iter_shapes(self._slide.shapes):
            applied |= _replace_in_shape(shp, replacements)
        return sorted(applied)

    # --- scoped replacement (repeated placeholders, per shape) -------------
    def fill_shapes(self, shape_map: dict[str, dict]) -> list[str]:
        """Fill placeholders scoped to specific NAMED shapes, so identical
        placeholder strings in sibling shapes (columns, stat cards, KPI blocks,
        card grids) get DIFFERENT values. `shape_map` maps a shape name to its
        own {find: replace} dict. Returns names of shapes not found."""
        missing = []
        for name, repl in shape_map.items():
            shp = self.find_shape(name)
            if shp is None:
                missing.append(name)
                continue
            _replace_in_shape(shp, repl)
        return missing

    def fill_runs(self, shape_name: str, texts: list) -> bool:
        """Assign new text to a shape's runs POSITIONALLY (run 0, run 1, …),
        preserving each run's formatting. Use to rewrite a mixed-emphasis line
        (e.g. a statement with one italicized word) where the emphasis lives on a
        specific run. A None/absent entry leaves that run untouched; passing ""
        blanks it. Returns True if the shape was found."""
        shp = self.find_shape(shape_name)
        if shp is None or not shp.has_text_frame:
            return False
        runs = [r for p in shp.text_frame.paragraphs for r in p.runs]
        for r, t in zip(runs, texts):
            if t is not None:
                r.text = t
        return True

    def fill_table(self, shape_name: str, cells) -> bool:
        """Fill a table's cells, preserving cell formatting. `cells` is either a
        2-D list of rows (row-major; None leaves a cell untouched) or a list of
        [row, col, text] triples for sparse edits. Returns True if found."""
        shp = self.find_shape(shape_name)
        if shp is None or not getattr(shp, "has_table", False):
            return False
        tbl = shp.table
        if cells and isinstance(cells[0], (list, tuple)) and len(cells[0]) == 3 \
                and isinstance(cells[0][0], int):
            for r, c, text in cells:
                _set_cell_text(tbl.cell(r, c), "" if text is None else str(text))
        else:
            for r, row in enumerate(cells):
                for c, text in enumerate(row):
                    if text is not None:
                        _set_cell_text(tbl.cell(r, c), str(text))
        return True

    # --- progress-divider (agenda reused as a section tracker) ------------
    def fill_divider(self, sections: list, active=None,
                     active_color=None, muted_color=None) -> bool:
        """Turn a duplicated agenda slide (canonical 24-28) into a progress
        divider: list every section, HIGHLIGHT the current one (ink + bold) and
        mute the rest (gray). `active` is the 0-based index of the current
        section (None = plain contents list, nothing highlighted). Reuse the
        SAME sections list on every divider so the deck reads as one consistent
        tracker. Returns True if section labels were found and filled."""
        act = _DIVIDER_ACTIVE if active_color is None else RGBColor.from_string(active_color)
        mut = _DIVIDER_MUTED if muted_color is None else RGBColor.from_string(muted_color)
        # Collect the section-label paragraphs in reading order (they read
        # "Section 01"…"Section 0N", whether in one list shape or separate ones).
        targets = []
        for shp in _iter_shapes(self._slide.shapes):
            if not shp.has_text_frame:
                continue
            for para in shp.text_frame.paragraphs:
                if _SECTION_RE.match("".join(r.text for r in para.runs)):
                    targets.append(para)
        if not targets:
            return False
        for i, para in enumerate(targets):
            if not para.runs:
                continue
            label = sections[i] if i < len(sections) else ""
            para.runs[0].text = label
            for extra in para.runs[1:]:
                extra.text = ""
            color = act if (active is not None and i == active) else mut
            para.runs[0].font.color.rgb = color
            if active is not None:
                para.runs[0].font.bold = (i == active)
        return True

    # --- image replacement -------------------------------------------------
    def replace_image(self, shape_name: str, image_path: str) -> bool:
        """Put an image into the named shape, keeping the shape's existing frame,
        size, crop, and position.

        Two cases, and only the first used to work:

        1. The shape already holds a photo (shape_type 13, PICTURE) — swap the
           blip relationship so the new image inherits the frame and crop.
        2. The shape is an EMPTY picture placeholder (shape_type 14,
           PLACEHOLDER, of type PICTURE) — fill it via insert_picture, which
           crops to the placeholder frame and leaves geometry untouched.

        Case 2 was missing, so any canonical slide whose image slot ships empty
        reported "shape not found" and rendered as a grey box. That is most of
        them: 152's slots ship with stock photos and worked, while 134's and
        85's ship empty and silently failed.
        """
        for shp in self._slide.shapes:
            if shp.name != shape_name:
                continue
            # case 1 — an existing picture: swap the blip, keep the frame
            if shp.shape_type == 13:  # PICTURE
                blip = shp._element.find(".//" + qn("a:blip"))
                if blip is None:
                    return False
                image_part, rId = self._slide.part.get_or_add_image_part(image_path)
                blip.set(qn("r:embed"), rId)
                _recrop(shp, image_path)
                return True
            # case 2 — an empty picture placeholder: insert into the frame
            if getattr(shp, "is_placeholder", False):
                try:
                    if "PICTURE" in str(shp.placeholder_format.type):
                        pic = shp.insert_picture(image_path)
                        _recrop(pic, image_path)
                        return True
                except (AttributeError, ValueError, KeyError):
                    return False
        return False


def _recrop(pic, image_path: str) -> None:
    """Re-derive a symmetric centre-fill crop for a newly placed image.

    Canonical picture slots ship carrying the crop of whatever stock photo the
    template designer framed there — slide 152's third slot, for instance, crops
    68% off the bottom and 18% off the left. Those values survive both a blip
    swap and insert_picture, so ANY replacement image inherited a stranger's
    framing and rendered as a mangled fragment. Recompute the crop from the new
    image's own aspect instead: nothing if the caller pre-fitted the image to the
    frame, otherwise an even centre crop on the overflowing axis. Position and
    size are untouched — only the crop, which belonged to the old photo.
    """
    try:
        fw, fh = pic.width, pic.height
        iw, ih = pic.image.size
    except (AttributeError, ValueError, TypeError):
        return
    if not (fw and fh and iw and ih):
        return
    frame_ar, img_ar = fw / fh, iw / ih
    if abs(img_ar - frame_ar) / frame_ar < 0.02:      # already fitted
        l = r = t = b = 0.0
    elif img_ar > frame_ar:                            # wider — trim the sides
        keep = frame_ar / img_ar
        l = r = (1.0 - keep) / 2.0
        t = b = 0.0
    else:                                              # taller — trim top/bottom
        keep = img_ar / frame_ar
        t = b = (1.0 - keep) / 2.0
        l = r = 0.0
    pic.crop_left, pic.crop_right = l, r
    pic.crop_top, pic.crop_bottom = t, b


class TemplateDeck:
    def __init__(self, template_path: str):
        self.template_path = Path(template_path)
        if not self.template_path.exists():
            raise FileNotFoundError(
                f"Canonical template not found: {self.template_path}. "
                "Slides must be duplicated from the canonical deck — there is no "
                "from-scratch fallback."
            )
        # Work inside a copy of the canonical deck so masters/layouts/theme/
        # fonts/media are all present and linked. The copy lives in a temp dir so
        # an interrupted run never litters the project folder.
        fd, tmp = tempfile.mkstemp(suffix=".work.pptx")
        os.close(fd)
        self._work = Path(tmp)
        shutil.copyfile(self.template_path, self._work)
        self.prs = Presentation(str(self._work))
        self._original_count = len(self.prs.slides._sldIdLst)
        self._used: list[Slide] = []

    def close(self):
        """Release the temporary working copy without touching the template."""
        work = getattr(self, "_work", None)
        if work is not None:
            try:
                work.unlink()
            except OSError:
                pass

    def __enter__(self):
        return self

    def __exit__(self, _exc_type, _exc, _tb):
        self.close()

    def __del__(self):
        self.close()

    # --- inspection -------------------------------------------------------
    @property
    def slide_count(self) -> int:
        return self._original_count

    def list_slide_text(self, number: int):
        """Return replaceable paragraph keys for canonical slide `number`.

        Empty paragraphs and shape-level newline joins are omitted because the
        replacement engine operates within runs and paragraphs, never across a
        paragraph boundary.
        """
        slide = self.prs.slides[number - 1]
        out = []
        for shp in _iter_shapes(slide.shapes):
            if not shp.has_text_frame:
                continue
            for para in shp.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs)
                if text:
                    out.append((shp.name, text))
        return out

    # --- duplication ------------------------------------------------------
    def use(self, number: int) -> Slide:
        """Duplicate canonical slide `number` (1-based) and return it for
        filling. The same template slide may be used more than once."""
        if not (1 <= number <= self._original_count):
            raise IndexError(
                f"Canonical slide {number} out of range 1..{self._original_count}"
            )
        src = self.prs.slides[number - 1]
        clone = self._clone(src)
        s = Slide(clone, number)
        self._used.append(s)
        return s

    def _clone(self, src_slide):
        # Use the SAME layout as the source slide so placeholders, theme color
        # mapping, and inherited formatting resolve identically.
        layout = src_slide.slide_layout
        new_slide = self.prs.slides.add_slide(layout)

        # Copy shapes INTO the new slide's existing spTree element (rather than
        # swapping the whole cSld) so python-pptx's cached shape collection
        # stays bound to the live element.
        new_spTree = new_slide.shapes._spTree
        for child in list(new_spTree):
            if child.tag in _SHAPE_TAG_SET:
                new_spTree.remove(child)  # drop the layout's placeholder shapes
        src_spTree = src_slide.shapes._spTree
        for child in src_spTree:
            if child.tag in _SHAPE_TAG_SET:
                new_spTree.append(copy.deepcopy(child))

        new_csld = new_slide._element.find(qn("p:cSld"))
        src_csld = src_slide._element.find(qn("p:cSld"))

        # Carry over a per-slide background fill, if the source has one.
        src_bg = src_csld.find(qn("p:bg"))
        if src_bg is not None:
            old_bg = new_csld.find(qn("p:bg"))
            if old_bg is not None:
                new_csld.remove(old_bg)
            new_csld.insert(0, copy.deepcopy(src_bg))  # bg precedes spTree

        # Carry over the source slide's color-map override.
        src_clr = src_slide._element.find(qn("p:clrMapOvr"))
        if src_clr is not None:
            new_sld = new_slide._element
            old_clr = new_sld.find(qn("p:clrMapOvr"))
            if old_clr is not None:
                new_sld.remove(old_clr)
            new_sld.append(copy.deepcopy(src_clr))

        # Rewrite relationship references so images/media resolve in the clone.
        _remap_rels(src_slide.part, new_slide.part, new_spTree)
        new_bg = new_csld.find(qn("p:bg"))
        if new_bg is not None:
            _remap_rels(src_slide.part, new_slide.part, new_bg)
        return new_slide

    # --- finalize ---------------------------------------------------------
    def save(self, out_path: str):
        """Drop the original canonical slides, leaving only the duplicated
        (and filled) clones in the order they were created, then write."""
        if not self._used:
            raise RuntimeError("No slides were used; nothing to save.")
        sld_id_lst = self.prs.slides._sldIdLst
        originals = list(sld_id_lst)[: self._original_count]
        for sld_id in originals:
            rId = sld_id.get(qn("r:id"))
            sld_id.getparent().remove(sld_id)
            try:
                self.prs.part.drop_rel(rId)
            except KeyError:
                pass
        self.prs.save(str(out_path))
        self.close()
        return out_path


__all__ = ["TemplateDeck", "Slide"]
