#!/usr/bin/env python3
"""build_from_plan.py — produce an Attentive deck by DUPLICATING canonical
template slides and filling them. This is the sanctioned entry point. It does
not generate any geometry; it only selects, duplicates, and fills.

Usage
-----
    python build_from_plan.py plan.json out.pptx
    python build_from_plan.py --inspect 24      # print canonical slide 24's text

    # A fit pre-flight runs automatically before every build. If any field
    # overflows its box (COMPRESS/SPLIT), a key doesn't match, or a slide ref is
    # out of range, the build is REFUSED so overflow never ships:
    python build_from_plan.py plan.json out.pptx --force    # build anyway
    python build_from_plan.py plan.json out.pptx --no-lint  # skip the gate

    # A post-build VERIFICATION then reads the finished .pptx (about a second)
    # and asserts what the pre-flight structurally cannot see: leftover canonical
    # dummy copy, repeated slots addressed out of reading order, images carrying
    # a previous photo's crop, unreplaced stock photos, and negative/roadmap
    # template variants. Errors exit non-zero — the file is still written so you
    # can open it and look.
    python build_from_plan.py plan.json out.pptx --verify-render  # + pixel checks
    python build_from_plan.py plan.json out.pptx --no-verify      # skip it

    # After filling each slide the builder auto-widens a short content-header box
    # so its title (and any subtitle) sits on ONE line, using only the free space
    # beside it — the sole geometry change allowed on a duplicated slide. It never
    # shrinks a box, crosses a neighbour, or touches display titles / agenda lists.
    python build_from_plan.py plan.json out.pptx --no-fit-titles  # leave headers raw
    #   or disable per step with  "fit_title": false

Plan format (JSON)
------------------
{
  "template": "Template—2026 NEW Attentive Company Deck Template.pptx",
  "slides": [
    {
      "template_slide": 1,                       # 1-based slide # in canonical deck
      "note": "cover",                           # optional, for your own tracing
      "text": {"Existing copy": "New copy"},     # WHOLE-slide find -> replace
      "images": {"Picture 3": "assets/hero.png"} # shape name -> new image path
    },
    {"template_slide": 24, "text": {"Old header": "New header"}}
  ]
}

Filling REPEATED placeholders (columns, stat trios, card grids, tables) — the
whole-slide `text` map replaces every copy of a string with the same value, so
sibling shapes that share a placeholder ("Title of section", identical lorem)
can't be filled differently that way. Address them per shape instead:

    "shapes": {                                  # find->replace scoped per shape
      "Google Shape;2021;p270": {"Title of section": "White-glove support",
                                 "Lorem ipsum …": "A dedicated strategist …"},
      "Google Shape;2022;p270": {"Title of section": "Identity + AI", …}
    },
    "runs":   {"Google Shape;1918;p259":          # positional run rewrite (keeps
               ["Your customers ", "", "phones", "", "", "", " …"]}, #  per-run emphasis)
    "tables": {"Google Shape;3394;p355":          # 2-D rows OR [row,col,text] list
               [["Functionality","Attentive","Klaviyo"], ["Support","✓","—"]]}

Reusing the agenda as a SECTION TRACKER (progress divider) — duplicate an agenda
slide (24-28) and highlight the current section:

    {"template_slide": 27, "note": "section: proof",
     "divider": {"sections": ["The opportunity","Why brands switch","One platform",
                              "Proof & results","Getting started","Next steps"],
                 "active": 3}}

`template_slide` is a reference to a physical canonical slide to DUPLICATE — not
a design spec to rebuild. Use --inspect to see a slide's exact placeholder text
before writing the maps. The same template slide may appear multiple times.
"""
import json
import sys
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).parent))
from template_deck import TemplateDeck
import preflight_lint
import antibland
import fit_titles  # the one sanctioned geometry change: one-line content headers

DEFAULT_TEMPLATE = "Template—2026 NEW Attentive Company Deck Template.pptx"


def _resolve_template(name: str, base_dir: Path | None = None) -> Path:
    here = Path(__file__).parent
    requested = Path(name).expanduser()
    if requested.is_absolute():
        candidates = [requested]
    elif name == DEFAULT_TEMPLATE:
        # The packaged canonical deck is authoritative. A same-named file in the
        # caller's working directory must never shadow it.
        candidates = [here / name, here.parent / name]
        if base_dir is not None:
            candidates.append(base_dir / requested)
        candidates.append(requested)
    else:
        candidates = []
        if base_dir is not None:
            candidates.append(base_dir / requested)
        candidates.extend([requested, here / requested, here.parent / requested])
    for cand in candidates:
        if cand.exists():
            return cand.resolve()
    hint = ""
    if (here / "TEMPLATE-REQUIRED.md").exists():
        hint = ("\n\nThis install ships without the canonical deck. Drop\n"
                f"  {name}\ninto\n  {here}\nkeeping the filename exactly as above. "
                "See TEMPLATE-REQUIRED.md in that folder.")
    raise FileNotFoundError(f"Canonical template not found: {name}{hint}")


def inspect(number: int, template: str = DEFAULT_TEMPLATE):
    deck = TemplateDeck(str(_resolve_template(template)))
    try:
        print(f"Canonical slide {number} of {deck.slide_count}:")
        for name, text in deck.list_slide_text(number):
            shown = text.replace("\n", " ⏎ ")
            print(f"  [{name}] {shown!r}")
    finally:
        deck.close()


def build(plan_path: str, out_path: str, fit_titles_enabled: bool = True):
    plan_file = Path(plan_path).expanduser().resolve()
    plan = json.loads(plan_file.read_text())
    plan_dir = plan_file.parent
    template = plan.get("template", DEFAULT_TEMPLATE)
    deck = TemplateDeck(str(_resolve_template(template, base_dir=plan_dir)))

    slide_w_in = deck.prs.slide_width / fit_titles.EMU_PER_IN
    print(f"Duplicating from: {template}  ({deck.slide_count} canonical slides)")
    for i, step in enumerate(plan["slides"], 1):
        num = step["template_slide"]
        slide = deck.use(num)
        note = step.get("note", "")
        applied = []
        # whole-slide find/replace
        if step.get("text"):
            applied = slide.replace_text(step["text"])
            missing = [k for k in step["text"] if k not in applied]
            if missing:
                print(
                    f"  ! slide {i}: {len(missing)} text key(s) not found on "
                    f"canonical slide {num}: {missing}. Run --inspect {num} to "
                    f"see the exact placeholder text."
                )
        # per-shape scoped fills (repeated placeholders: columns, stat cards, …)
        if step.get("shapes"):
            gone = slide.fill_shapes(step["shapes"])
            for name in gone:
                print(f"  ! slide {i}: shape {name!r} not found on canonical #{num}.")
            applied += [k for m in step["shapes"].values() for k in m]
        # positional per-run rewrite (mixed emphasis lines)
        for shape_name, texts in (step.get("runs") or {}).items():
            if not slide.fill_runs(shape_name, texts):
                print(f"  ! slide {i}: shape {shape_name!r} (runs) not on canonical #{num}.")
            else:
                applied += [t for t in texts if t]
        # table cell fills
        for shape_name, cells in (step.get("tables") or {}).items():
            if not slide.fill_table(shape_name, cells):
                print(f"  ! slide {i}: table {shape_name!r} not found on canonical #{num}.")
            else:
                applied.append(shape_name)
        # progress divider (reused agenda with the active section highlighted)
        div = step.get("divider")
        if div:
            ok = slide.fill_divider(div.get("sections", []), div.get("active"),
                                    div.get("active_color"), div.get("muted_color"))
            if not ok:
                print(f"  ! slide {i}: no section labels found to fill on canonical #{num}.")
            else:
                applied += div.get("sections", [])
        for shape_name, img in (step.get("images") or {}).items():
            image_path = Path(img).expanduser()
            if not image_path.is_absolute():
                image_path = (plan_dir / image_path).resolve()
            ok = slide.replace_image(shape_name, str(image_path))
            if not ok:
                print(
                    f"  ! slide {i}: image shape {shape_name!r} not found on "
                    f"canonical slide {num}."
                )
        # One-line content-header auto-fit — the ONLY geometry change we allow on
        # a duplicated slide: widen a short content-header box into the free space
        # beside it so the title/subtitle sits on one line (never shrinks, never
        # crosses a neighbour, never touches display titles). Skip globally with
        # --no-fit-titles or per step with "fit_title": false.
        widened = []
        if fit_titles_enabled and step.get("fit_title", True):
            widened = fit_titles.widen_titles(slide._slide, slide_w_in)
        print(
            f"  slide {i:>2}: duplicated canonical #{num}"
            f"{' (' + note + ')' if note else ''}"
            f" — filled {len(applied)} block(s)"
        )
        for nm, old_in, new_in in widened:
            print(f"          · title {nm!r} → 1 line ({old_in}→{new_in} in)")

    deck.save(out_path)
    print(f"\nWrote {out_path} — {len(plan['slides'])} duplicated template slides.")
    print(
        "QA: every slide above traces to a canonical slide number. Render and "
        "eyeball before shipping."
    )
    return out_path


def _fresh_facts(template_path: Path) -> bool:
    """Use the packaged facts without mutating the installed skill.

    Facts are generated and validated when Vanessa publishes the skill. Installed
    organization skills may be read-only, so a deck build must never rewrite a
    resource in place.
    """
    facts = Path(__file__).parent / "slide_facts.json"
    if not facts.exists():
        print("(packaged slide_facts.json is missing; semantic trap checks skipped)")
        return False
    bundled_template = (Path(__file__).parent / DEFAULT_TEMPLATE).resolve()
    if template_path.resolve() != bundled_template:
        print("(custom template selected; packaged semantic trap checks skipped)")
        return False
    return True


def _post_verify(out_path: str, plan_path: str, template_path: Path,
                 do_render: bool) -> int:
    """Assert on the FINISHED deck. Returns the number of errors.

    The pre-flight reasons about the plan; this reads the artefact. It exists
    because on the AI Guild rebrand seven of eleven defects were invisible until
    something was rendered and inspected by hand — repeated slots filled in
    shape-id rather than reading order, images inheriting the previous photo's
    crop, panels rendering as pale boxes, canonical stock photos shipping.
    """
    import verify_render

    print("\n" + "=" * 68)
    print("POST-BUILD VERIFICATION")
    have_facts = _fresh_facts(template_path)
    prs = Presentation(out_path)
    rep = verify_render.Report()

    canonical_strings = set()
    inv = Path(__file__).parent / "inventory.json"
    if inv.exists():
        for entries in json.loads(inv.read_text()).values():
            for _, v in entries:
                for line in str(v).replace("\x0b", "\n").split("\n"):
                    line = line.strip()
                    if len(line) > 3:
                        canonical_strings.add(line)
    verify_render.check_dummy_text(prs, canonical_strings, rep)

    if have_facts:
        facts = json.loads((Path(__file__).parent / "slide_facts.json").read_text())
        plan = json.loads(Path(plan_path).read_text())
        verify_render.check_slot_order(plan, facts, rep)
        verify_render.check_semantics(plan, facts, rep)

    verify_render.check_images(
        prs, verify_render.canonical_image_hashes(str(template_path)), rep)

    if do_render:
        pages = verify_render.render_pages(out_path)
        if pages is None:
            rep.add("WARN", "render", None, "could not render (soffice missing?)")
        else:
            verify_render.check_rendered(prs, pages, rep)

    rep.print()
    n_e, n_w = len(rep.errors), len(rep.warns)
    print(f"\n{n_e} error(s), {n_w} warning(s)")
    if n_e:
        print("⛔ VERIFICATION FAILED — the deck was written, open it and look at "
              "the slides above. Re-run with --no-verify to ignore.")
    else:
        print("✓ verification passed")
    if not do_render:
        print("  (pixel checks skipped — add --verify-render for those)")
    return n_e


def _gate(plan_path: str, force: bool, no_lint: bool):
    """Run the pre-flight before building. Refuse to build on overflow /
    unmatched keys / out-of-range slides — AND on a monotonous template
    selection or an inconsistent transition system — unless --force is given."""
    if no_lint:
        print("(pre-flight skipped: --no-lint)\n")
        return
    report = preflight_lint.lint_plan(plan_path)
    preflight_lint.print_report(report)
    print()
    ok = report.get("ok_to_build", report["clean"])
    if ok or force:
        if not ok:
            print("⚠ Building anyway (--force): the issues above will ship.\n")
        return
    reasons = []
    if not report["clean"]:
        reasons.append("fit problems (overflow / unmatched keys / out-of-range)")
    if not report.get("variety", {}).get("ok", True):
        reasons.append("a monotonous template selection (see MONOTONY above)")
    if not report.get("transitions", {}).get("ok", True):
        reasons.append("an inconsistent transition system")
    if not report.get("vividness", {}).get("ok", True) and antibland.GATE_ON_VIVIDNESS:
        reasons.append("a bland template selection (see ANTI-BLAND above)")
    print("⛔ Refusing to build — the plan has " + "; ".join(reasons) + ".")
    print("   Fix the flagged copy or diversify the templates (a denser/richer")
    print("   family, a stat/card/comparison/quote/timeline slide, or a split);")
    print("   lift any BLAND slide onto a stat/KPI/chart/statement/quote template;")
    print("   or re-run with --force to build anyway, or --no-lint to skip.")
    sys.exit(1)


def main():
    args = sys.argv[1:]
    if len(args) >= 2 and args[0] == "--inspect":
        inspect(int(args[1]))
        return
    force = "--force" in args
    no_lint = "--no-lint" in args
    fit_titles_enabled = "--no-fit-titles" not in args
    no_verify = "--no-verify" in args
    verify_render_too = "--verify-render" in args
    pos = [a for a in args if not a.startswith("--")]
    if len(pos) < 2:
        print(__doc__)
        sys.exit(1)
    _gate(pos[0], force, no_lint)
    build(pos[0], pos[1], fit_titles_enabled=fit_titles_enabled)
    if no_verify:
        return
    plan = json.loads(Path(pos[0]).read_text())
    plan_file = Path(pos[0]).expanduser().resolve()
    template = _resolve_template(
        plan.get("template", DEFAULT_TEMPLATE), base_dir=plan_file.parent
    )
    # --force covers the PLAN gate; it deliberately does not silence verification,
    # which is about the artefact that just got written.
    if _post_verify(pos[1], pos[0], template, verify_render_too):
        sys.exit(1)


if __name__ == "__main__":
    main()
